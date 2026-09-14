from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
ROOT=Path(__file__).resolve().parents[3]; OUT=ROOT/'.build/msup-demo/deliverables'; SHOTS=ROOT/'.build/msup-demo/screenshots-v2';OUT.mkdir(exist_ok=True)
font='/System/Library/Fonts/Hiragino Sans GB.ttc'
im=Image.new('RGB',(1600,720),'white');d=ImageDraw.Draw(im)
f=ImageFont.truetype(font,32);small=ImageFont.truetype(font,24)
rows=[('接入与数据','司空连接器与历史文件  →  项目  设备  时空资产','连接专业平台，汇聚业务素材'),('智能能力','算法服务产生证据  →  智能体给出建议','可替换的感知能力与有依据的业务建议'),('任务编排','触发器  →  观察  →  检测  →  研判  →  工单  →  报告','Task 发布版本与 Run 执行记录'),('业务协作','人工复核  →  业务工单  →  报告与证据回溯','围绕同一份依据持续协作')]
for i,(a,b,c) in enumerate(rows):
 y=20+i*170;d.rounded_rectangle((20,y,1580,y+145),radius=10,fill='#f1f5f9',outline='#cbd5e1',width=2);d.text((45,y+18),a,font=f,fill='#0f172a');d.text((300,y+18),b,font=f,fill='#0f172a');d.text((300,y+80),c,font=small,fill='#475569')
im.save(OUT/'architecture.png')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Arial';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB');st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(9)
doc.styles['Title'].font.size=Pt(25);doc.styles['Title'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 1'].paragraph_format.space_after=Pt(13)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.color.rgb=RGBColor.from_string('475569')
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
text=(ROOT/'docs/demo/project-report-draft.md').read_text();parts=re.split(r'^## ',text,flags=re.M);head=parts[0]
doc.add_paragraph('AeroSight 低空巡检\n感知与决策协同平台',style='Title')
doc.add_paragraph('项目报告    高校组    2026年9月',style='Subtitle')
doc.add_paragraph('参赛队伍：一只汤圆叮\n所属学校：大连理工大学')
doc.add_paragraph('赛题三 AI+时空智能赛道\n面向低空飞行的感知 决策一体化智能体开发')
doc.add_picture(str(OUT/'AeroSight封面.png'),width=Inches(7))
doc.add_page_break()
imgs={1:('architecture.png','图 1 平台架构与模块协同'),2:('observation.png','图 2 数据资产与观察范围'),3:('detection.png','图 3 视觉检测与原图关联'),4:('assessment.png','图 4 智能体建议与证据引用'),5:('task.png','图 5 任务定义与 YAML 编辑入口'),6:('issue.png','图 6 工单协作与证据关联'),8:('report.png','图 7 报告汇总与关联入口')}
for i,part in enumerate(parts[1:]):
 title,body=part.split('\n',1)
 if i>0:doc.add_page_break()
 doc.add_heading(title,level=1)
 paras=[x.strip() for x in body.strip().split('\n\n') if x.strip()]
 for p in paras:doc.add_paragraph(p)
 if i in imgs:
  filename,caption=imgs[i];path=OUT/filename if i==1 else SHOTS/filename
  doc.add_picture(str(path),width=Inches(7));doc.paragraphs[-1].paragraph_format.keep_with_next=True;doc.add_paragraph(caption,style='Caption')
for el in list(doc.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for p in doc.paragraphs:
 for el in list(p._p.iter(qn('w:pBdr'))): el.getparent().remove(el)
doc.styles['Subtitle'].font.italic=False
doc.styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
doc.core_properties.title='AeroSight 低空巡检感知与决策协同平台项目报告';doc.core_properties.author='大连理工大学 一只汤圆叮'
doc.save(OUT/'AeroSight项目报告.docx')
print(OUT/'AeroSight项目报告.docx')
