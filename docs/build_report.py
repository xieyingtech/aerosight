from __future__ import annotations

from pathlib import Path
from datetime import date
from urllib.parse import urlparse

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "report_assets"
OUT = ROOT / "AeroSight空天一体化智能感知平台技术方案报告.docx"

NAVY = "163A5F"
BLUE = "2166A5"
SKY = "EAF3FA"
TEAL = "0B7285"
ORANGE = "D97706"
GOLD = "F4B740"
INK = "1F2937"
MUTED = "64748B"
LIGHT = "F4F6F9"
LINE = "D7E0E8"
WHITE = "FFFFFF"
GREEN = "2F855A"
RED = "A13D3D"

SOURCES = [
    ("DJI Cloud API：产品介绍", "https://developer.dji.com/doc/cloud-api-tutorial/en/overview/product-introduction.html"),
    ("DJI Cloud API：功能与教程总览", "https://developer.dji.com/doc/cloud-api-tutorial/en/tutorial-map.html"),
    ("DJI Dock 2：官方规格", "https://enterprise.dji.com/dock-2/specs"),
    ("DJI FlightHub 2：官方产品介绍", "https://enterprise.dji.com/flighthub-2"),
    ("DJI Dock 2 发布说明", "https://enterprise.dji.com/news/detail/dji-dock-2-release"),
    ("Real-Time Object Detection Based on UAV Remote Sensing: A Systematic Literature Review", "https://doi.org/10.3390/drones7100620"),
    ("Using UAV Remote Sensing and a Monitoring Information System to Enhance the Management of Unauthorized Structures", "https://doi.org/10.3390/app9224954"),
    ("YOLO-UB Based Detection and Identification of Illegal Structures in the Ancient City", "https://doi.org/10.3233/FAIA230864"),
    ("JointNet4BCD: Semi-Supervised Building Change Detection", "https://doi.org/10.3390/rs16234569"),
    ("Building Extraction from Remote Sensing Images with Deep Learning: A Survey", "https://doi.org/10.1016/j.cviu.2024.104253"),
    ("《无人驾驶航空器飞行管理暂行条例》", "https://www.caac.gov.cn/XXGK/XXGK/FLFG/202401/t20240115_222642.html"),
    ("民航局关于民用无人驾驶航空器监管服务有关事宜的公告", "https://www.caac.gov.cn/XXGK/XXGK/TZTG/202312/t20231231_222550.html"),
    ("《中华人民共和国个人信息保护法》", "https://www.cac.gov.cn/2021-08/20/c_1631050028355286.htm"),
    ("《中华人民共和国数据安全法》", "https://www.npc.gov.cn/npc/c2/c30834/202106/t20210610_311888.html"),
]


def font_path(preferred: list[str]) -> str:
    windows_fonts = Path("C:/Windows/Fonts")
    for name in preferred:
        p = windows_fonts / name
        if p.exists():
            return str(p)
    raise FileNotFoundError("No suitable font found")


FONT_REG = font_path(["msyh.ttc", "msyh.ttf", "simhei.ttf", "arial.ttf"])
FONT_BOLD = font_path(["msyhbd.ttc", "msyhbd.ttf", "simhei.ttf", "arialbd.ttf"])


def img_font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def rounded(draw, xy, radius=20, fill=WHITE, outline=LINE, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=f"#{fill}", outline=f"#{outline}", width=width)


def centered_text(draw, box, text, font, fill=INK, spacing=6):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=spacing)
    w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=font, fill=f"#{fill}", align="center", spacing=spacing)


def arrow(draw, start, end, color=BLUE, width=7):
    draw.line([start, end], fill=f"#{color}", width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - sign * 18, y2 - 11), (x2 - sign * 18, y2 + 11)]
    else:
        sign = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 11, y2 - sign * 18), (x2 + 11, y2 - sign * 18)]
    draw.polygon(pts, fill=f"#{color}")


def build_architecture_png(path: Path):
    im = Image.new("RGB", (1800, 1040), "#F8FAFC")
    d = ImageDraw.Draw(im)
    d.text((70, 48), "AeroSight 空天一体化智能感知平台总体架构", font=img_font(48, True), fill=f"#{NAVY}")
    d.text((72, 112), "设备接入、边缘感知、时空数据、核心智能体、业务闭环", font=img_font(25), fill=f"#{MUTED}")

    layers = [
        ("空天感知层", "无人机 / 大疆机场\n可见光、热红外、倾斜摄影\nGNSS/RTK、气象、设备遥测", "EAF3FA", BLUE),
        ("边缘智能层", "YOLO 检测与轻量跟踪\n关键帧筛选、脱敏与压缩\n离线缓存、事件优先上报", "E8F5F2", TEAL),
        ("平台与时空数据层", "DJI Cloud API / MQTT / HTTPS\n对象存储、PostGIS、时序库\n影像/轨迹/事件统一索引", "FFF4E6", ORANGE),
        ("核心智能体层", "任务理解、知识检索、工具调用\n多时相对比、规划约束核验\n置信度融合、风险分级与解释", "EEF2FF", NAVY),
        ("业务执行层", "问题工单、人工复核、处置协同\n报告生成、复飞核验、统计看板\n模型回流与持续优化", "F3E8FF", "7C3AED"),
    ]
    y = 205
    for i, (title, body, fill, accent) in enumerate(layers):
        x1, x2 = 120, 1680
        rounded(d, (x1, y, x2, y + 128), radius=22, fill=fill, outline=accent, width=3)
        d.rounded_rectangle((x1 + 20, y + 22, x1 + 275, y + 106), radius=16, fill=f"#{accent}")
        centered_text(d, (x1 + 20, y + 22, x1 + 275, y + 106), title, img_font(29, True), WHITE)
        d.multiline_text((x1 + 330, y + 25), body, font=img_font(24), fill=f"#{INK}", spacing=9)
        if i < len(layers) - 1:
            arrow(d, (900, y + 128), (900, y + 158), color=MUTED, width=5)
        y += 158

    d.rounded_rectangle((1450, 46, 1690, 126), radius=18, fill=f"#{NAVY}")
    centered_text(d, (1450, 46, 1690, 126), "数据—空间—智能—任务", img_font(22, True), WHITE)
    im.save(path, quality=95)


def build_loop_png(path: Path):
    im = Image.new("RGB", (1800, 860), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.text((65, 40), "城市违建巡查示例：从自动巡逻到销号复核", font=img_font(46, True), fill=f"#{NAVY}")
    d.text((68, 103), "YOLO 负责快速发现，核心智能体负责结合时空、规划与业务证据做综合判断", font=img_font(24), fill=f"#{MUTED}")
    items = [
        ("1", "任务生成", "定时/事件触发\n空域、天气、设备校验"),
        ("2", "自动巡逻", "机场起飞\n按航线采集影像与遥测"),
        ("3", "端侧初筛", "YOLO 检测\n关键帧、目标框、置信度"),
        ("4", "Agent 研判", "多时相变化 + GIS/规划许可\n知识与规则融合、风险分级"),
        ("5", "处置与报告", "工单、证据包、人工复核\n形成规范报告并流转"),
        ("6", "复飞与学习", "复查整改状态、销号\n误报/漏报回流训练"),
    ]
    x = 70
    y = 220
    box_w, box_h, gap = 245, 265, 42
    accents = [BLUE, TEAL, ORANGE, NAVY, "7C3AED", GREEN]
    for i, (num, title, body) in enumerate(items):
        rounded(d, (x, y, x + box_w, y + box_h), radius=24, fill=LIGHT, outline=accents[i], width=3)
        d.ellipse((x + 18, y + 18, x + 76, y + 76), fill=f"#{accents[i]}")
        centered_text(d, (x + 18, y + 18, x + 76, y + 76), num, img_font(26, True), WHITE)
        d.text((x + 22, y + 98), title, font=img_font(29, True), fill=f"#{accents[i]}")
        d.multiline_text((x + 22, y + 150), body, font=img_font(20), fill=f"#{INK}", spacing=9)
        if i < len(items) - 1:
            arrow(d, (x + box_w + 4, y + box_h / 2), (x + box_w + gap - 5, y + box_h / 2), color=MUTED, width=5)
        x += box_w + gap
    arrow(d, (1515, 530), (1515, 675), color=GREEN, width=6)
    arrow(d, (1515, 675), (215, 675), color=GREEN, width=6)
    arrow(d, (215, 675), (215, 530), color=GREEN, width=6)
    d.rounded_rectangle((605, 638, 1125, 716), radius=18, fill="#E7F5EC", outline=f"#{GREEN}", width=3)
    centered_text(d, (605, 638, 1125, 716), "闭环：结果反馈 → 数据治理 → 模型/规则持续优化", img_font(23, True), GREEN)
    im.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(w))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.42)
    sec.footer_distance = Inches(0.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 14, MUTED, 0, 8),
        ("Heading 1", 17, BLUE, 16, 8),
        ("Heading 2", 13.5, NAVY, 11, 5),
        ("Heading 3", 11.5, TEAL, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = name != "Subtitle"
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.1

    for list_name in ("List Bullet", "List Number"):
        st = styles[list_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.5)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.194)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.208


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    r2 = paragraph.add_run(" 页")
    set_font(r2, 8.5, color=MUTED)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def p(doc, text="", bold_lead=None, align=None, after=None, keep=False):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = para.add_run(bold_lead)
        set_font(r1, bold=True, color=NAVY)
        r2 = para.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = para.add_run(text)
        set_font(r)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    r = para.add_run(text)
    set_font(r)
    return para


def numbered(doc, text):
    para = doc.add_paragraph(style="List Number")
    r = para.add_run(text)
    set_font(r)
    return para


def add_heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    r = para.add_run(text)
    set_font(r, bold=True, color=BLUE if level == 1 else (NAVY if level == 2 else TEAL))
    return para


def add_callout(doc, label, text, fill=SKY, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:color"), accent)
    borders.append(left)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    r = para.add_run(f"{label}  ")
    set_font(r, 10.5, True, accent)
    r = para.add_run(text)
    set_font(r, 10.5, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=9.1):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(h)
        set_font(run, font_size, True, WHITE)
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F8FAFC")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.12
            run = para.add_run(str(val))
            set_font(run, font_size, False, INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path, caption, width=6.45):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    picture = para.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    set_font(r, 8.5, False, MUTED)


def setup_headers(doc):
    sec = doc.sections[0]
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("AEROSIGHT  |  空天一体化智能感知平台")
    set_font(r, 8.5, True, NAVY)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), LINE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    fp = sec.footer.paragraphs[0]
    add_page_field(fp)


def cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = kicker.add_run("AEROSIGHT 技术方案简报")
    set_font(r, 12, True, ORANGE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("空天一体化智能感知平台")
    set_font(r, 30, True, NAVY)
    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("以无人机自动巡逻与城市违建治理闭环为示范场景")
    set_font(r, 15, False, MUTED)
    doc.add_paragraph()
    add_callout(doc, "核心主张", "平台的价值不止于“端侧跑 YOLO”，而在于把无人机、机场、时空数据、模型、知识、规则、智能体与城市治理流程连接成可感知、可解释、可决策、可执行、可追溯的完整系统。", fill="EEF5FA", accent=NAVY)
    for _ in range(4):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("基于 AeroSight 当前项目骨架、现有无人机/大疆机场与违建数据集\n调研日期：2026 年 8 月 21 日")
    set_font(r, 10.5, False, MUTED)
    doc.add_page_break()


def executive_summary(doc):
    add_heading(doc, "摘要", 1)
    p(doc, "AeroSight 具备建设“空天一体化智能感知平台”的良好起点：已有无人机与大疆机场可承担常态化数据采集，已有较大规模违建数据集可用于模型训练与持续迭代，当前软件项目则已形成项目、设备、设备能力、任务、执行记录、素材、问题、智能体会话等基础对象。建议以城市违建巡查作为第一条可验证闭环，但平台从一开始按多设备、多任务、多模型、多智能体和多业务场景设计。")
    add_callout(doc, "建议结论", "先做“单机场、单片区、单一违建类型”的可度量试点，用 8–12 周跑通任务生成—自动飞行—端侧初筛—Agent 研判—人工复核—工单报告—复飞销号；在数据、接口和安全边界稳定后，再扩展到道路病害、河道、工地、应急等场景。", fill="FFF6E8", accent=ORANGE)
    add_table(doc,
        ["维度", "判断", "对本项目的含义"],
        [
            ("产品定位", "不是单点识别算法，而是时空智能业务操作系统", "以任务闭环和证据链为产品中心，算法作为可插拔能力"),
            ("技术路线", "边缘快速筛查 + 云/中心综合研判 + 人在回路", "弱网可运行、关键证据优先、重大结论需可解释和复核"),
            ("落地优势", "设备、机场、数据集均已具备", "可直接进入真实场景基线测试，而非停留在仿真或公开数据"),
            ("首要难点", "变化检测、地理配准、规划数据核验比单帧目标检测更关键", "YOLO 发现“像不像”，Agent 判断“是不是、为何、如何处置”"),
            ("建设策略", "复用 DJI Cloud API，强化 AeroSight 的业务编排与智能体层", "避免重复制造成熟的飞控能力，把投入集中到差异化闭环"),
        ], [1500, 2600, 5260])


def current_state(doc):
    add_heading(doc, "1. 项目基础与边界", 1)
    add_heading(doc, "1.1 当前 AeroSight 已有的软件骨架", 2)
    p(doc, "代码仓库显示，AeroSight 当前采用 Go + Gin + PostgreSQL 的后端与 Next.js + MapLibre 的前端，已具备团队/项目权限、地图总览以及设备、任务、素材、问题和智能体页面。数据库对象之间已经能表达从任务执行到证据素材、问题事件、智能体会话的基本关联。")
    add_table(doc,
        ["现有对象/页面", "已具备内容", "建议补强"],
        [
            ("devices / device_capabilities", "设备状态、最后在线时间、能力代码与约束", "大疆机场/无人机 Thing Model 映射、健康度、气象与载荷状态"),
            ("tasks / task_runs", "触发方式、能力要求、目标选择、脚本、输入输出快照", "航线版本、前置安全校验、重试/补飞、SLA 与审批门"),
            ("assets", "设备、任务、问题关联的影像/文件及元数据", "地理范围、姿态/镜头、哈希、模型版本、证据链与生命周期"),
            ("issues / issue_events", "问题状态、优先级、时间线和关联对象", "疑似违建案件字段、空间几何、人工复核、处置部门与销号标准"),
            ("agents / sessions / messages", "Agent 配置、会话、消息、工具调用和 token 记录", "工具注册表、知识来源、策略版本、审计轨迹、结构化结论"),
            ("MapLibre 项目地图", "基础地图容器", "航线、机场、禁限飞区、目标框/面、热力图与多时相对比"),
        ], [2100, 3000, 4260], font_size=8.8)
    add_heading(doc, "1.2 本报告的系统边界", 2)
    bullet(doc, "空中端：无人机、机场、载荷、飞控与航线执行；优先复用大疆现成能力。")
    bullet(doc, "边缘端：可部署在无人机伴随计算单元、机场侧计算盒或近场边缘服务器；承担实时检测、关键帧筛选、脱敏和缓存。")
    bullet(doc, "平台端：设备接入、任务编排、时空数据治理、模型服务、智能体编排、业务工单和报告。")
    bullet(doc, "人机协同：AI 生成“疑似事件和建议”，执法或业务人员完成法定核验与最终处置。")


def landscape(doc):
    add_heading(doc, "2. 既有系统与论文的主要启示", 1)
    add_heading(doc, "2.1 工业系统：成熟飞控云平台已覆盖“飞起来”", 2)
    p(doc, "DJI FlightHub 2 已提供远程操控、智能调度、航线管理、数据分析和第三方集成；DJI Cloud API 使用 MQTT、HTTPS、WebSocket 等通用协议，将设备抽象为物联网 Thing Model，并覆盖设备管理、航线任务、直播流、媒体上传和态势感知。Dock 2 官方规格也明确支持 FlightHub 2、第三方云平台接入与外接边缘计算通信。因此 AeroSight 更合理的定位，是在成熟设备控制之上建立跨设备的时空数据底座、智能体决策与城市治理闭环。")
    add_table(doc,
        ["方案类型", "强项", "不足/机会", "AeroSight 取舍"],
        [
            ("DJI FlightHub 2 + Dock", "设备兼容、远程飞行、航线、直播与自动化流程成熟", "面向通用无人机运营，业务规则和跨系统治理需二次集成", "作为设备与飞行底座；AeroSight 承担业务和 Agent 层"),
            ("第三方 Drone-in-a-box 平台", "自动值守、远程运维、多站点调度", "通常绑定硬件/生态，国内合规和既有设备适配成本不同", "借鉴远程运维与安全门设计，不作为首期替换目标"),
            ("传统 WebGIS 违建系统", "多期影像、DSM、矢量数据库和案件管理相对成熟", "自动化与实时性有限，人工判读成本高", "吸收时空数据库、变化对比与案件证据链"),
            ("纯视觉算法 Demo", "训练迭代快，便于展示检测效果", "缺少任务、设备、空间、法规和处置闭环", "仅作为边缘感知插件，不定义整个平台"),
        ], [1800, 2600, 2600, 2360], font_size=8.4)
    add_heading(doc, "2.2 学术研究：违建识别应从单帧检测走向时空证据融合", 2)
    p(doc, "相关研究大致分为三条路线：一是 UAV 影像上的 YOLO/CNN 目标检测或分类；二是正射影像、DSM/点云与多时相影像的建筑变化检测；三是结合 WebGIS、地籍/规划数据和人工复核的管理系统。系统综述指出，UAV 实时目标检测的完整实现必须同时评估精度、速度、时延与能耗，边缘计算是主流部署范式。违建监管研究则表明，仅凭外观分类不足以认定违法，需把高度变化、建筑边界、多期影像以及业务许可数据纳入证据链。")
    add_callout(doc, "研究启示", "首期可以用 YOLO 快速发现彩钢棚、屋顶加建、施工材料和新生结构等可见线索；真正降低误报的关键，是同一地块多时相配准、建筑轮廓/高度变化、规划许可和历史案件的联合判断。", fill="E8F5F2", accent=TEAL)


def architecture(doc, arch_path):
    add_heading(doc, "3. 总体架构：数据—空间—智能—任务", 1)
    p(doc, "参考图强调的不是某个模型，而是围绕低空智能基础设施、数字化管控、数据空间治理和大模型/智能体形成系统化能力。AeroSight 建议采用“五层一横向”架构：五层负责业务数据流动，一条横向安全与运维体系贯穿全链路。")
    add_figure(doc, arch_path, "图 1  AeroSight 总体架构（本报告设计）")
    add_heading(doc, "3.1 五层能力", 2)
    add_table(doc,
        ["层级", "核心组件", "关键输出"],
        [
            ("空天感知层", "无人机、机场、相机/热红外、RTK、气象、遥测", "标准化设备状态、航迹、原始影像与环境信息"),
            ("边缘智能层", "YOLO、跟踪、关键帧、裁剪/脱敏、离线缓存", "疑似目标事件包；弱网下继续执行并择机上报"),
            ("平台与时空数据层", "Cloud API 网关、对象存储、PostGIS、时序库、特征库", "影像—位置—时间—设备—任务—案件统一索引"),
            ("核心智能体层", "任务理解、知识库、工具调用、时空计算、规划与反思", "带证据、置信度、风险等级和下一步建议的结构化结论"),
            ("业务执行层", "问题中心、工单、复核、报告、通知、复飞与销号", "可追溯处置结果与学习反馈"),
        ], [1600, 4200, 3560], font_size=8.8)
    add_heading(doc, "3.2 横向保障", 2)
    bullet(doc, "安全：设备身份、双向认证、最小权限、密钥轮换、存储加密、下载水印和操作审计。")
    bullet(doc, "可靠性：设备心跳、任务状态机、幂等事件、断点续传、降级策略、人工接管与应急返航。")
    bullet(doc, "模型治理：数据集版本、模型版本、阈值策略、灰度发布、漂移监测和一键回滚。")
    bullet(doc, "可观测性：飞行、链路、推理、Agent、工单全链路 trace；核心指标统一看板。")


def workflow(doc, loop_path):
    add_heading(doc, "4. 示例闭环：无人机自动巡逻查找城市违建", 1)
    add_figure(doc, loop_path, "图 2  城市违建巡查端到端闭环（本报告设计）")
    add_heading(doc, "4.1 事件数据包：端侧只上传“足够的证据”", 2)
    p(doc, "端侧 YOLO 不直接输出“违法”结论，而是生成 suspected_object 事件。建议事件包包含：任务与设备 ID、UTC 时间、经纬度/高度/姿态、航线与相机参数、目标类别、边界框/分割掩膜、置信度、连续帧轨迹、全景关键帧、局部裁剪、模型版本、原始文件哈希和边缘设备健康状态。低置信事件仅保留摘要，高风险或连续出现事件优先上传完整证据。")
    add_heading(doc, "4.2 核心智能体如何综合判断", 2)
    numbered(doc, "理解任务：识别巡查区域、违建类型、时间窗口、飞行与业务约束。")
    numbered(doc, "校验证据：检查时间戳、定位、姿态、图像质量、模型版本和证据完整性。")
    numbered(doc, "调用工具：查询历史影像、同地块上期结果、规划许可/地籍信息、已有案件与 GIS 图层。")
    numbered(doc, "开展时空分析：完成影像配准、目标地理定位、变化面积/高度估计、邻近关系与趋势判断。")
    numbered(doc, "执行规则与模型融合：规则用于硬约束和法规口径，视觉/变化模型用于感知，大模型用于解释与流程编排。")
    numbered(doc, "输出结构化结论：疑似类型、空间范围、证据清单、置信度、风险等级、解释、缺失信息和下一步动作。")
    numbered(doc, "人在回路：低风险抽检，中高风险人工复核；确认后生成问题工单与标准报告。")
    numbered(doc, "复飞销号：整改期限到达后自动生成复查任务；结果回写案件并进入训练/评估集。")
    add_heading(doc, "4.3 建议的判断分层", 2)
    add_table(doc,
        ["等级", "机器结论", "处理策略", "示例"],
        [
            ("L0", "无明显异常", "留存摘要，按比例抽检", "检测框短暂出现、图像模糊或与历史一致"),
            ("L1", "待观察变化", "进入复飞队列或调整视角补拍", "目标可见但定位/遮挡导致证据不足"),
            ("L2", "疑似违建", "自动建问题，人工复核后派单", "新增屋顶结构且无可关联许可"),
            ("L3", "高风险疑似", "即时告警、优先复核与现场核查", "大面积快速施工、占用消防通道或重大安全风险"),
        ], [1100, 2450, 2900, 2910], font_size=8.7)
    add_callout(doc, "责任边界", "平台只能输出“疑似违建/风险线索”，不能替代主管部门的法定认定。所有证据、规则版本、人工操作与报告版本应可追溯。", fill="FDEEEE", accent=RED)


def agent_and_data(doc):
    add_heading(doc, "5. 核心智能体与数据/模型体系", 1)
    add_heading(doc, "5.1 核心智能体的六类能力", 2)
    add_table(doc,
        ["能力", "作用", "建议工具"],
        [
            ("任务理解", "把自然语言或业务规则转成目标、约束和验收条件", "任务模板、规则解析、结构化输出 Schema"),
            ("知识检索", "引用法规、规划口径、案例和设备手册", "版本化知识库、来源引用、有效期检查"),
            ("工具调用", "连接设备、GIS、数据库、模型、工单和报告服务", "受控 Tool Registry、权限与超时/重试"),
            ("时空计算", "定位、范围、距离、轨迹、变化、遮挡与邻接分析", "PostGIS、影像配准、DSM/点云、地图服务"),
            ("规划执行", "拆解任务、选择设备/航线、补拍与复飞", "状态机、约束求解、人审门、补偿动作"),
            ("反思与评估", "检查证据缺口、冲突和不确定性", "置信度校准、规则一致性、离线评测与回放"),
        ], [1700, 3800, 3860], font_size=8.7)
    add_heading(doc, "5.2 模型不应只有一个 YOLO", 2)
    p(doc, "推荐形成“实时模型 + 离线精判模型 + 规则/知识”的组合。端侧采用轻量检测/分割模型满足帧率和功耗；云端对疑似区域运行更大分辨率的检测、分割与多时相变化模型；当条件允许时引入正射影像、DSM/点云和建筑轮廓矢量。大模型不直接阅读所有视频，而是消费结构化事件、精选图像、时空查询结果和规则结论。")
    add_table(doc,
        ["模型/能力", "部署位置", "目标指标", "说明"],
        [
            ("轻量 YOLO/分割", "边缘端", "实时性、召回率、功耗", "高召回初筛；按场景动态阈值"),
            ("目标跟踪与关键帧", "边缘端", "重复事件压缩率、轨迹稳定性", "避免同一目标反复上报"),
            ("高分辨率精判", "中心端", "精确率、边界质量", "对疑似目标二次推理"),
            ("多时相变化检测", "中心端", "变化 IoU/F1、配准鲁棒性", "识别新增、扩建、拆除和持续施工"),
            ("DSM/点云分析", "中心端", "高度误差、体量估计", "增强对屋顶加建和体量变化的证据"),
            ("Agent 融合与解释", "中心端", "可追溯率、证据完整率、人工采纳率", "输出结构化结论，不越过法定认定边界"),
        ], [2100, 1550, 2450, 3260], font_size=8.5)
    add_heading(doc, "5.3 数据飞轮", 2)
    p(doc, "现有违建数据集是项目的关键资产，但需要补齐“真实飞行域”的数据闭环。建议建立 raw（原始）、curated（清洗标注）、golden（固定评测）、hard-case（难例）四层数据集；严格划分地理区域和时间，避免相邻帧泄漏到训练/验证集造成虚高。每次模型发布必须绑定数据版本、训练配置、阈值和评测报告。")
    bullet(doc, "优先标注：违建类型、建筑/地块几何、遮挡/光照、飞行高度与视角、是否变化、最终复核结论。")
    bullet(doc, "主动学习：优先回流低置信、模型冲突、人工改判和新区域样本。")
    bullet(doc, "评测分层：离线数据集、历史任务回放、影子模式、有限灰度和正式运行。")


def implementation(doc):
    add_heading(doc, "6. 与现有代码衔接的实施建议", 1)
    p(doc, "现有数据库结构适合作为控制面和业务对象骨架，但生产系统还需要补充空间数据、设备事件、模型推理与证据链。建议保持现有 Go API、Next.js 和 PostgreSQL 技术路线，逐步引入 PostGIS、对象存储、消息总线与模型服务，不做一次性重写。")
    add_table(doc,
        ["新增/扩展模块", "落到现有对象", "首期最小实现"],
        [
            ("DJI 接入网关", "devices、device_capabilities、task_runs", "Cloud API MQTT/HTTPS；设备影子、航线下发、状态回调、媒体上传"),
            ("时空事件中心", "assets、issues、issue_events", "PostGIS geometry；事件包 schema；地图展示目标点/面与航迹"),
            ("模型服务", "assets.metadata_json、task_runs.output_snapshot_json", "端侧推理协议、模型注册、版本与阈值、批量回放"),
            ("Agent 工具层", "agents、agent_sessions、agent_messages", "GIS 查询、影像对比、许可查询、工单创建、报告生成五类工具"),
            ("案件与报告", "issues、issue_links、issue_events", "复核状态机、证据清单、报告模板、复飞任务与销号"),
            ("运营看板", "项目地图和各列表页", "设备在线率、任务成功率、疑似事件、复核结论、覆盖率与趋势"),
        ], [2000, 2950, 4410], font_size=8.6)
    add_heading(doc, "6.1 建议的关键状态机", 2)
    p(doc, "任务状态建议扩展为 queued → precheck → dispatched → executing → uploading → analyzing → review_required → completed / failed / cancelled；问题状态建议扩展为 suspected → triaged → verified / rejected → assigned → rectifying → reinspection → closed。每次状态变化都写入不可变事件日志，便于审计和重放。")
    add_heading(doc, "6.2 首期接口原则", 2)
    bullet(doc, "设备命令与业务命令分离：Agent 不直接调用飞控底层，必须通过带策略校验的任务服务。")
    bullet(doc, "同步查询与异步任务分离：长时飞行、媒体处理和模型推理均由事件驱动，API 只返回任务 ID。")
    bullet(doc, "结构化结果优先：Agent 输出 JSON Schema，同时生成面向人的解释和报告。")
    bullet(doc, "所有自动动作可撤销或人工接管：航线下发、复飞、告警和案件创建均设权限与审批门。")


def pilot(doc):
    add_heading(doc, "7. 真实试点方案与验收指标", 1)
    add_heading(doc, "7.1 8–12 周试点路线", 2)
    add_table(doc,
        ["阶段", "周期", "主要工作", "阶段产物"],
        [
            ("P0 基线", "第 1–2 周", "设备接入；选定片区/违建类型；梳理数据与飞行合规；建立人工真值", "航线、事件 Schema、基线模型与安全清单"),
            ("P1 影子运行", "第 3–5 周", "自动飞行与端侧初筛；Agent 只给建议不自动建案；全量人工对照", "误报/漏报分析、阈值和难例集"),
            ("P2 人机闭环", "第 6–8 周", "中高置信自动建问题；人工复核；生成标准报告；安排复飞", "完整案件链、报告模板、闭环看板"),
            ("P3 稳定性", "第 9–12 周", "多天气/多时段回归；断网、失败、补飞和人工接管演练", "试点评估报告、扩面与采购决策"),
        ], [1300, 1100, 4300, 2660], font_size=8.5)
    add_heading(doc, "7.2 验收指标建议", 2)
    add_table(doc,
        ["指标域", "核心指标", "首期建议目标/口径"],
        [
            ("飞行与设备", "任务成功率、设备在线率、异常返航/人工接管", "任务成功率 ≥95%；所有失败原因可归类、可追溯"),
            ("边缘推理", "事件召回率、端侧时延、重复压缩率", "以真实片区人工真值为准；先保召回，再通过云端降误报"),
            ("综合研判", "案件级 precision/recall、证据完整率、人工采纳率", "按违建类型、区域、天气与高度分层统计，不只报总体 mAP"),
            ("业务效率", "每平方公里巡查成本、单事件复核时长、报告生成时长", "与人工巡查基线对照；报告自动生成后人工只做核验和签发"),
            ("系统闭环", "从发现到建单/销号时长、可追溯率", "关键动作审计覆盖率 100%；证据链无孤儿对象"),
            ("安全合规", "越界、失联、未授权访问和隐私事件", "重大安全事件为 0；每项演练有记录与整改"),
        ], [1600, 3500, 4260], font_size=8.7)
    add_callout(doc, "评测提醒", "公开数据上的 mAP 不能替代现场验收。必须用本地航线、真实高度/视角、不同季节与天气、相邻建筑遮挡和同地块多时相数据建立 golden set，并按案件级而非单帧级计算效果。", fill="FFF6E8", accent=ORANGE)


def risks_and_roadmap(doc):
    add_heading(doc, "8. 风险、合规与规模化路线", 1)
    add_heading(doc, "8.1 关键风险控制", 2)
    add_table(doc,
        ["风险", "表现", "控制措施"],
        [
            ("飞行安全与空域", "常态巡逻涉及固定空域、远程/网络链路与人口密集区", "实名登记、空域/计划流程、起飞前检查、电子围栏、人工接管和应急预案"),
            ("隐私与数据安全", "航拍可能采集人脸、车牌、住宅窗户和敏感地理信息", "目的限定、最小采集、边缘脱敏、分级权限、留存期限、访问审计与安全评估"),
            ("误报/错判", "单帧外观无法证明是否合法", "多时相/GIS/许可融合；证据不足降级；最终法定认定由人员完成"),
            ("数据漂移", "季节、天气、城市风貌和飞行参数变化", "分层监测、难例回流、区域化阈值、灰度发布和回滚"),
            ("供应商锁定", "设备接口与云能力依赖单一生态", "内部统一设备抽象与事件 Schema；保留多厂商适配层"),
            ("Agent 不可控", "幻觉、越权工具调用或循环执行", "白名单工具、结构化输出、预算/超时、策略引擎、人审门和完整审计"),
        ], [1700, 3300, 4360], font_size=8.5)
    p(doc, "合规方面，《无人驾驶航空器飞行管理暂行条例》自 2024 年 1 月 1 日起施行，要求飞行组织者承担安全主体责任，并对登记、识别信息、飞行申请/备案和应急处置等作出规定；民航局 UOM 平台承担实名登记、适飞空域查询和相关申请服务。城市航拍数据还应遵循个人信息保护、数据安全以及测绘地理信息管理等要求。具体试点应由属地主管部门、飞行运营方和法律/安全人员联合确认边界。")
    add_heading(doc, "8.2 规模化路线", 2)
    bullet(doc, "从单机场到多机场：增加覆盖规划、任务冲突消解、备降/接力、站点健康和统一运维。")
    bullet(doc, "从单模型到模型市场：检测、分割、变化检测、热异常、三维重建按场景编排。")
    bullet(doc, "从单智能体到多智能体：任务智能体、感知研判智能体、合规智能体、报告智能体分工协作，核心编排器统一权限与状态。")
    bullet(doc, "从违建到城市治理：复用同一数据—空间—智能—任务底座，扩展工地、河道、道路、应急、园区和基础设施巡检。")
    add_heading(doc, "9. 结论", 1)
    p(doc, "AeroSight 最具价值的方向，是成为连接低空设备与城市治理业务的时空智能操作平台。违建巡查适合作为首个标杆闭环：它同时要求自动飞行、边缘视觉、地理定位、多时相变化、规则知识、人工复核、案件处置和报告生成，能够真实检验平台是否“可感知、可解释、可决策、可执行”。凭借现有无人机/大疆机场、违建数据集和已成形的软件对象模型，项目已经具备进入现场试点的条件。下一步应以可度量的单片区试点替代大而全建设，在闭环数据中逐步把算法优势沉淀为系统壁垒。")


def references(doc):
    add_heading(doc, "参考资料", 1)
    p(doc, "以下资料用于本报告的系统与技术调研。产品能力以官方文档为准；论文指标受数据集与实验设置影响，不直接等同于本项目现场效果。")
    for idx, (title, url) in enumerate(SOURCES, 1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.25)
        para.paragraph_format.first_line_indent = Inches(-0.25)
        para.paragraph_format.space_after = Pt(5)
        r = para.add_run(f"[{idx}] {title}. ")
        set_font(r, 9.2, False, INK)
        add_hyperlink(para, url, url)
    add_heading(doc, "项目内依据", 2)
    p(doc, "AeroSight README、db/schema.sql、internal/store/models.go、internal/store/queries.go，以及 web/app/(app)/projects/[id] 下的地图、设备、任务、素材、问题和智能体页面（检视日期：2026 年 8 月 21 日）。")


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    arch = ASSET_DIR / "architecture.png"
    loop = ASSET_DIR / "closed_loop.png"
    build_architecture_png(arch)
    build_loop_png(loop)

    doc = Document()
    style_document(doc)
    setup_headers(doc)
    cover(doc)
    executive_summary(doc)
    current_state(doc)
    landscape(doc)
    architecture(doc, arch)
    workflow(doc, loop)
    agent_and_data(doc)
    implementation(doc)
    pilot(doc)
    risks_and_roadmap(doc)
    references(doc)

    props = doc.core_properties
    props.title = "AeroSight 空天一体化智能感知平台技术方案报告"
    props.subject = "无人机自动巡逻、边缘智能、核心智能体与城市违建治理闭环"
    props.author = "AeroSight 项目组"
    props.keywords = "AeroSight, 无人机, 大疆机场, YOLO, 智能体, 违建巡查, 时空智能"
    props.comments = "基于项目代码、用户提供参考图与公开资料形成的技术方案。"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
